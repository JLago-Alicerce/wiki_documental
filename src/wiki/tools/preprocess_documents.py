from __future__ import annotations

from pathlib import Path
import re
import docx
from docx.document import Document
from docx.text.paragraph import Paragraph

try:
    from pdf2docx import Converter
except Exception:
    Converter = None  # type: ignore

try:
    from pdf2image import convert_from_path
except Exception:
    convert_from_path = None  # type: ignore

try:
    from pytesseract import image_to_string
except Exception:
    image_to_string = None  # type: ignore

from yaml import safe_load

from .docx_utils import clean_docx_styles
from . import pdf_utils

# Re-export Converter for backwards compatibility with tests
Converter = pdf_utils.Converter

# --- Heurísticas de detección de encabezados ---
HEADING_1_MIN_FONT_SIZE = 14  # pt
HEADING_2_MIN_FONT_SIZE = 12

def is_likely_heading(p: Paragraph) -> int:
    text = p.text.strip()
    if not text:
        return 0
    return 0  # Placeholder logic, implement actual heuristics

def _remove_paragraphs(paragraphs: list[Paragraph]) -> None:
    for p in paragraphs:
        element = p._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)

def _discard_sections(doc: Document, cfg: dict) -> bool:
    opts = cfg.get("options", {}).get("discard_sections", {})
    paragraphs = list(doc.paragraphs)

    def contains(text: str, words: list[str]) -> bool:
        t = text.lower()
        return any(w in t for w in words)

    modified = False

    if opts.get("portada", False) and paragraphs:
        first_block = []
        year_found = False
        logo_found = False
        for p in paragraphs:
            if p.style.name.startswith("Heading") or len(first_block) >= 5:
                break
            first_block.append(p)
            if re.search(r"\b\d{4}\b", p.text):
                year_found = True
            if "logo" in p.text.lower():
                logo_found = True
        if first_block and year_found and logo_found:
            _remove_paragraphs(first_block)
            print("🗑️ Portada descartada")
            paragraphs = list(doc.paragraphs)
            modified = True

    if opts.get("indice", False) and paragraphs:
        idx_paragraphs = [
            p for p in paragraphs if contains(p.text, ["indice", "índice", "contenido", "contents"]) or "......" in p.text
        ]
        if idx_paragraphs:
            _remove_paragraphs(idx_paragraphs)
            print("🗑️ Índice descartado")
            paragraphs = list(doc.paragraphs)
            modified = True

    if opts.get("contraportada", False) and paragraphs:
        tail_paragraphs = paragraphs[-5:]
        to_remove = [p for p in tail_paragraphs if contains(p.text, ["agradec", "logo", "fecha"])]
        if to_remove:
            _remove_paragraphs(to_remove)
            print("🗑️ Contraportada descartada")
            modified = True

    return modified

def clean_docx_styles(doc_path: Path, output_path: Path, cfg: dict | None = None) -> bool:
    try:
        doc: Document = docx.Document(doc_path)
    except Exception as e:
        print(f"❌ No se pudo abrir {doc_path.name}: {e}")
        return False

    changes_made = False
    if cfg:
        changes_made = _discard_sections(doc, cfg)

    for p in doc.paragraphs:
        if p.style.name.startswith('Heading'):
            continue
        heading_level = is_likely_heading(p)
        if heading_level > 0:
            print(f"  → '{p.text[:50].strip()}' → Heading {heading_level}")
            p.style = f'Heading {heading_level}'
            changes_made = True

    if changes_made:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)

    return changes_made

def process_pdf_with_converter(pdf_path: Path, output_docx_path: Path) -> None:
    pdf_utils.Converter = Converter
    pdf_utils.convert_pdf_to_docx(pdf_path, output_docx_path)

def process_pdf_with_ocr(pdf_path: Path, output_dir: Path, *, include_images: bool = True) -> None:
    if convert_from_path is None or image_to_string is None:
        raise ImportError("pdf2image and pytesseract are required for OCR")
    print(f"🧠 Ejecutando OCR para {pdf_path.name}")
    pages = convert_from_path(str(pdf_path))

    output_dir.mkdir(parents=True, exist_ok=True)
    assets_dir = output_dir / "assets"
    assets_dir.mkdir(parents=True, exist_ok=True)

    docx_path = output_dir / f"{pdf_path.stem}.docx"
    document = docx.Document()

    for idx, img in enumerate(pages, start=1):
        text = image_to_string(img)
        document.add_heading(f"Página {idx}", level=2)
        document.add_paragraph(text.strip())

        img_name = f"pagina_{idx:02}.png"
        img_path = assets_dir / img_name
        img.save(img_path)
        if include_images:
            document.add_picture(str(img_path))

    document.save(docx_path)

def batch_process_directory(input_dir: Path, output_dir: Path, cfg: dict | None = None):
    print(f"\n📂 Procesando documentos desde: {input_dir}")

    ocr_enabled = False
    include_images = True
    ocr_dir = output_dir.parent / "ocr_outputs"
    if cfg:
        ocr_cfg = cfg.get("options", {}).get("ocr", False)
        if isinstance(ocr_cfg, bool):
            ocr_enabled = ocr_cfg
        else:
            ocr_enabled = ocr_cfg.get("enabled", False)
            include_images = ocr_cfg.get("include_images", True)
        ocr_dir = Path(cfg["paths"].get("work", output_dir.parent)) / "ocr_outputs"

    for file in sorted(input_dir.glob("*")):
        print(f"\n📝 {file.name}")
        output_file = output_dir / (file.stem + ".docx")

        if file.suffix.lower() == ".docx":
            clean_docx_styles(file, output_file, cfg or {})

        elif file.suffix.lower() == ".pdf":
            try:
                if ocr_enabled:
                    raise RuntimeError("Forzando OCR por configuración")
                process_pdf_with_converter(file, output_file)
                clean_docx_styles(output_file, output_file, cfg or {})
            except Exception as e:
                print(f"❌ Error al convertir {file.name}: {e}")
                pdf_output_dir = ocr_dir / file.stem
                process_pdf_with_ocr(file, pdf_output_dir, include_images=include_images)

    print("\n✅ Limpieza completada.")

if __name__ == "__main__":
    cfg_path = Path("config.yaml")
    cfg = safe_load(cfg_path.read_text(encoding="utf-8"))

    input_dir = Path(cfg["paths"]["originals"])
    output_dir = Path(cfg["paths"]["cleaned"])
    output_dir.mkdir(parents=True, exist_ok=True)
    batch_process_directory(input_dir, output_dir, cfg)
