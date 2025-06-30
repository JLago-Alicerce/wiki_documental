from __future__ import annotations

from pathlib import Path

from docx import Document
import logging

STYLE_TO_LEVEL = {"Nav_Tit_1": 1, "Nav_Tit_2": 2, "Nav_Tit_3": 3}
from ..tools.docx_utils import get_safe_heading_style

from .normalization_utils import (
    average_font_size,
    detect_heading_by_fallback,
    detect_heading_by_heuristics,
    detect_heading_by_style,
    should_use_heuristics,
)


def _remove_toc_paragraphs(document: Document) -> None:
    """Delete paragraphs that belong to a table of contents."""
    for paragraph in list(document.paragraphs):
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.startswith("TOC") or "......" in paragraph.text:
            element = paragraph._element
            element.getparent().remove(element)


def normalize_styles(doc_path: Path, out_path: Path, cfg: dict | None = None) -> None:
    """Normalize visual styles to structured heading styles in a DOCX file."""
    if cfg is None:  # pragma: no cover - used by CLI
        from ..config import cfg as default_cfg

        cfg = default_cfg

    document = Document(str(doc_path))
    _remove_toc_paragraphs(document)

    use_heuristics = should_use_heuristics(doc_path, cfg)
    use_fallback = cfg.get("options", {}).get("use_heuristic_headings", False)
    avg_size = average_font_size(document) if use_fallback else 0.0

    for paragraph in document.paragraphs:
        level = detect_heading_by_style(paragraph)

        if level == 0:
            style_name = paragraph.style.name if paragraph.style else ""
            if style_name in STYLE_TO_LEVEL:
                level = STYLE_TO_LEVEL[style_name]

        if level == 0 and use_heuristics:
            level = detect_heading_by_heuristics(paragraph)

        if level == 0 and use_fallback:
            level = detect_heading_by_fallback(paragraph, avg_size)
            if level:
                logging.warning(
                    "Fallback heading detection used for paragraph: %s",
                    paragraph.text[:50],
                )

        if level:
            safe_style = get_safe_heading_style(document, level)
            paragraph.style = safe_style

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))
