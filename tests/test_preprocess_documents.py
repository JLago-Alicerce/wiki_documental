from pathlib import Path
from docx import Document
from docx.shared import Pt
from typer.testing import CliRunner

from wiki.cli import app

runner = CliRunner()


def _create_dirty_docx(path: Path) -> None:
    doc = Document()
    run = doc.add_paragraph().add_run("Title")
    run.bold = True
    run.font.size = Pt(16)
    doc.add_paragraph("Body text")
    doc.save(path)


def test_cli_preprocess_docs(tmp_path, monkeypatch):
    orig = tmp_path / "orig"
    work = tmp_path / "work"
    orig.mkdir()
    work.mkdir()

    doc_in = orig / "sample.docx"
    _create_dirty_docx(doc_in)

    pdf_in = orig / "sample_pdf.pdf"
    pdf_in.write_bytes(b"%PDF-1.4\n%%EOF")

    class DummyConverter:
        def __init__(self, src):
            self.src = src

        def convert(self, dest, start=0, end=None):
            doc = Document()
            run = doc.add_paragraph().add_run("PDF Title")
            run.bold = True
            run.font.size = Pt(14)
            doc.save(dest)

        def close(self):
            pass

    monkeypatch.setattr("wiki.tools.preprocess_documents.Converter", DummyConverter)
    paths = {
        "originals": orig,
        "cleaned": work / "cleaned",
        "to_process": work / "to_process",
    }
    monkeypatch.setattr("wiki.cli.cfg", {"paths": paths})

    result = runner.invoke(app, ["preprocess-docs"])
    assert result.exit_code == 0

    cleaned = paths["cleaned"] / "sample.docx"
    assert cleaned.exists()
    doc = Document(cleaned)
    assert doc.paragraphs[0].style.name.startswith("Heading")

    pdf_cleaned = paths["cleaned"] / "sample_pdf.docx"
    assert pdf_cleaned.exists()
    pdf_doc = Document(pdf_cleaned)
    assert pdf_doc.paragraphs[0].style.name == "Heading 1"

    # Cleaned files should be automatically copied for processing
    doc_copied = paths["to_process"] / "sample.docx"
    pdf_copied = paths["to_process"] / "sample_pdf.docx"
    assert doc_copied.exists()
    assert pdf_copied.exists()


def test_discard_sections(tmp_path, monkeypatch):
    orig = tmp_path / "orig"
    work = tmp_path / "work"
    orig.mkdir()
    work.mkdir()

    doc_in = orig / "sample.docx"
    doc = Document()
    doc.add_paragraph("MANUAL TECNICO")
    doc.add_paragraph("Fecha: 2024")
    doc.add_paragraph("LOGO")
    doc.add_paragraph("Contenido")
    doc.add_paragraph("Intro ...... 1")
    doc.add_paragraph("Introduccion", style="Heading 1")
    doc.add_paragraph("Texto real")
    doc.add_paragraph("Agradecimientos")
    doc.save(doc_in)

    cfg = {
        "paths": {"originals": orig, "cleaned": work / "cleaned", "to_process": work / "to_process"},
        "options": {"discard_sections": {"portada": True, "indice": True, "contraportada": True}},
    }
    monkeypatch.setattr("wiki.cli.cfg", cfg)

    result = runner.invoke(app, ["preprocess-docs"])
    assert result.exit_code == 0

    cleaned = cfg["paths"]["cleaned"] / "sample.docx"
    doc = Document(cleaned)
    texts = [p.text for p in doc.paragraphs]
    assert "MANUAL TECNICO" not in texts
    assert not any("Contenido" in t for t in texts)
    assert "Agradecimientos" not in texts
    assert "Texto real" in texts
