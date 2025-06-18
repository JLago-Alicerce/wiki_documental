from pathlib import Path
from docx import Document
from docx.shared import Pt

from wiki.tools.clean_docx import clean_docx_styles
from wiki.cli import app
from typer.testing import CliRunner

runner = CliRunner()


def _create_dirty_docx(path: Path) -> None:
    doc = Document()
    run = doc.add_paragraph().add_run("Title")
    run.bold = True
    run.font.size = Pt(16)
    doc.add_paragraph("Body text")
    doc.save(path)


def test_clean_docx_styles(tmp_path):
    dirty = tmp_path / "sample.docx"
    cleaned = tmp_path / "cleaned.docx"
    _create_dirty_docx(dirty)
    clean_docx_styles(dirty, cleaned)
    doc = Document(cleaned)
    assert doc.paragraphs[0].style.name == "Heading 1"


def test_cli_clean_docx(tmp_path, monkeypatch):
    orig = tmp_path / "orig"
    work = tmp_path / "work"
    orig.mkdir()
    work.mkdir()
    doc_in = orig / "sample.docx"
    _create_dirty_docx(doc_in)
    paths = {"originals": orig, "cleaned": work / "cleaned", "work": work}
    monkeypatch.setattr("wiki.cli.cfg", {"paths": paths})
    result = runner.invoke(app, ["clean-docx"])
    assert result.exit_code == 0
    cleaned = paths["cleaned"] / "sample.docx"
    assert cleaned.exists()
    doc = Document(cleaned)
    assert doc.paragraphs[0].style.name == "Heading 1"

